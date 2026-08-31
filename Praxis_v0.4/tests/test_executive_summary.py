"""Tests for llm_assist.draft_executive_summary - added after a real gap
was flagged: the executive summary of a data-driven report was a fixed
template regardless of whether an LLM was configured, producing the same
thin, generic-feeling paragraph even for a rich analysis (500 observations,
imputation, outlier treatment, a computed treatment effect). The fix
follows the same pattern already used elsewhere in this file (draft_
recommendations, draft_document_body): the LLM only *interprets* numbers
DataAnalysisAgent already computed - it never computes anything itself,
and never sees/produces a fact it wasn't given.
"""
from unittest.mock import patch

from src.services import llm_assist


class _FakeClient:
    available = True

    def __init__(self, capture=None, response="Interprétation rédigée à partir des faits fournis."):
        self._capture = capture
        self._response = response

    def complete(self, system, user, max_tokens=None):
        if self._capture is not None:
            self._capture["user"] = user
            self._capture["system"] = system
        return self._response

    def complete_json(self, system, user, max_tokens=None):
        # These tests are about the executive summary specifically;
        # draft_recommendations (also called by DocumentAgent) uses this
        # and gracefully falls back to the scaffold when it returns None.
        return None


def test_draft_executive_summary_is_grounded_in_real_facts_only():
    captured = {}
    with patch.object(llm_assist, "LLMClient", lambda: _FakeClient(captured)):
        result = llm_assist.draft_executive_summary(
            title="Évaluation Agri-Borgou", task_type="rapport_evaluation",
            raw_request="Évaluer l'impact du programme", objective=None,
            n_observations=500, n_anomalies=12, strategy="B",
            enrichment={"n_imputations": 3, "n_outliers_traites": 2},
            treatment_effect={
                "ate": 40129.5, "p_value": 0.0001, "significant_at_5pct": True,
                "outcome": "revenu_post - revenu_pre", "ci_95": [30000, 50000],
                "method": "Différence-en-différences (OLS)",
            },
        )
    assert result == "Interprétation rédigée à partir des faits fournis."
    assert "500 observations analysées" in captured["user"]
    assert "40,129.50" in captured["user"]
    assert "0.0001" in captured["user"]


def test_draft_executive_summary_none_without_llm():
    assert llm_assist.draft_executive_summary(
        title="t", task_type="analyse_donnees", raw_request="r", objective=None,
        n_observations=100, n_anomalies=0, strategy="B",
    ) is None  # no LLM configured in tests - honest fallback, not a crash


def test_draft_executive_summary_none_when_nothing_to_summarize():
    """No facts at all - nothing to ground an interpretation in, so this
    must not fabricate one even with an LLM configured (and must not even
    call it)."""
    class _AssertNoCall:
        available = True
        def complete(self, *a, **k):
            raise AssertionError("must not call the LLM with nothing to summarize")

    with patch.object(llm_assist, "LLMClient", _AssertNoCall):
        result = llm_assist.draft_executive_summary(
            title="t", task_type="analyse_donnees", raw_request="r", objective=None,
            n_observations=None, n_anomalies=None, strategy=None,
        )
    assert result is None


def test_presentation_agent_leads_with_llm_headline_when_available(tmp_path, pilot_xlsx):
    """Same fix, same reasoning, for the pptx 'Résultats clés' slide -
    the pasted diagnostic's other concrete complaint ('2 diapositives...
    une structure prédéfinie')."""
    from src.agents import DataAnalysisAgent, PresentationAgent
    from pptx import Presentation

    clean_result = DataAnalysisAgent().execute({
        "step": {"type": "nettoyage"}, "previous_outputs": {},
        "data_sources": [{"path": pilot_xlsx["path"]}], "strategy": "B",
        "work_dir": str(tmp_path / "clean"),
    })
    previous_outputs = {
        a["role"]: {"file_path": a["file_path"], "metadata": a["metadata"]}
        for a in clean_result["artifacts"]
    }
    base_context = {
        "title": "Rapport pilote", "task_type": "analyse_donnees",
        "raw_request": "Analyser l'enquête pilote", "objective": None,
        "data_sources": [], "previous_outputs": previous_outputs,
    }

    with patch.object(llm_assist, "LLMClient", lambda: _FakeClient(response="Synthèse pptx rédigée par l'IA.")):
        result = PresentationAgent().execute({**base_context, "work_dir": str(tmp_path / "pptx_with_llm")})

    prs = Presentation(next(a["file_path"] for a in result["artifacts"] if a["role"] == "presentation_pptx"))
    all_text = []
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                all_text.append(shape.text_frame.text)
    assert any("Synthèse pptx rédigée par l'IA." in t for t in all_text)


def test_document_agent_uses_llm_summary_when_available(tmp_path, pilot_xlsx):
    """End-to-end: with an LLM configured (faked), the docx's Résumé
    exécutif contains the LLM's interpretation, not the fixed template -
    and without one, it contains the template exactly as before."""
    from src.agents import DataAnalysisAgent, DocumentAgent
    from docx import Document

    clean_result = DataAnalysisAgent().execute({
        "step": {"type": "nettoyage"}, "previous_outputs": {},
        "data_sources": [{"path": pilot_xlsx["path"]}], "strategy": "B",
        "work_dir": str(tmp_path / "clean"),
    })
    previous_outputs = {
        a["role"]: {"file_path": a["file_path"], "metadata": a["metadata"]}
        for a in clean_result["artifacts"]
    }

    base_context = {
        "title": "Rapport pilote", "task_type": "analyse_donnees",
        "raw_request": "Analyser l'enquête pilote", "objective": None,
        "data_sources": [], "previous_outputs": previous_outputs,
    }

    # Without LLM: template.
    result_no_llm = DocumentAgent().execute({**base_context, "work_dir": str(tmp_path / "no_llm")})
    doc_no_llm = Document(next(a["file_path"] for a in result_no_llm["artifacts"] if a["role"] == "report_docx"))
    text_no_llm = "\n".join(p.text for p in doc_no_llm.paragraphs)
    assert "Cette analyse porte sur" in text_no_llm  # the fixed template's own wording

    # With LLM (faked): drafted interpretation instead.
    with patch.object(llm_assist, "LLMClient", lambda: _FakeClient(response="Synthèse rédigée par l'IA pour ce rapport.")):
        result_llm = DocumentAgent().execute({**base_context, "work_dir": str(tmp_path / "with_llm")})
    doc_llm = Document(next(a["file_path"] for a in result_llm["artifacts"] if a["role"] == "report_docx"))
    text_llm = "\n".join(p.text for p in doc_llm.paragraphs)
    assert "Synthèse rédigée par l'IA pour ce rapport." in text_llm
    assert "Cette analyse porte sur" not in text_llm
