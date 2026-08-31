"""Every TaskType must produce a real, working plan and deliverable -
these tests exist because REPONSE_AO, RAPPORT_EVALUATION, PLANIFICATION
and AUTRE used to fall through to a plan containing a step type
("execution") that matches no agent, silently reaching 'deliverable'
status with zero artifacts produced. See docs/REFONTE_v0.4.md."""
from src.services.task_service import TaskService
from src.services.knowledge_service import KnowledgeService
from src.services.execution_engine import ExecutionEngine
from src.models import TaskType, KBContentType


def _run_task(db_session, task_type, raw_request, dims, hard_constraints=None, data_path=None):
    ts = TaskService(db_session)
    task = ts.create_task(title=f"Test {task_type.value}", raw_request=raw_request, task_type=task_type)
    if hard_constraints:
        task.hard_constraints = hard_constraints
        db_session.commit()
    if data_path:
        task.data_sources = [{"path": data_path, "original_name": "data.xlsx"}]
        db_session.commit()
    ts.update_task_readiness(task.id, dims)
    plan = ts.propose_plan(task.id)["plan"]
    ts.validate_plan(plan.id, approved=True)
    result = ExecutionEngine(db_session).execute_plan(task.id)
    return task, plan, result


def test_every_plan_step_type_maps_to_a_real_agent(db_session):
    """The actual bug: 'execution' matched no keyword in
    Orchestrator.AGENT_MAPPING and silently resolved to a no-op agent."""
    from src.services.orchestrator import Orchestrator

    ts = TaskService(db_session)
    known_keywords = list(Orchestrator.AGENT_MAPPING.keys())
    for task_type in TaskType:
        task = ts.create_task(title="x", raw_request="x", task_type=task_type)
        steps = ts._generate_basic_plan_steps(task)
        for step in steps:
            assert any(kw in step["type"] for kw in known_keywords), (
                f"{task_type}: step type '{step['type']}' matches no agent"
            )


def test_reponse_ao_produces_real_deliverable_with_compliance_section(db_session, tmp_path):
    task, plan, result = _run_task(
        db_session, TaskType.REPONSE_AO, "Répondre à l'appel d'offres de rénovation d'école",
        {"objectif": 0.9}, hard_constraints=["Délai de remise: 30 jours", "Certification ISO 9001 requise"],
    )
    assert result["final_status"] == "deliverable"
    assert len(result["deliverables"]) == 1

    from docx import Document
    from src.services.artifact_service import ArtifactService
    artifact = ArtifactService(db_session).get_artifact(result["deliverables"][0]["artifact_id"])
    doc = Document(artifact.file_ref)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Conformité" in text
    assert "Délai de remise: 30 jours" in text


def test_planification_produces_real_deliverable(db_session):
    task, plan, result = _run_task(
        db_session, TaskType.PLANIFICATION, "Planifier la mise en œuvre du projet WASH sur 12 mois",
        {"objectif": 0.9},
    )
    assert result["final_status"] == "deliverable"
    assert len(result["artifacts"]) >= 1


def test_autre_produces_real_deliverable(db_session):
    task, plan, result = _run_task(
        db_session, TaskType.AUTRE, "Rédiger une note de service sur le télétravail", {"objectif": 0.9},
    )
    assert result["final_status"] == "deliverable"
    assert len(result["deliverables"]) == 1


def test_rapport_evaluation_without_data_uses_content_driven_plan(db_session):
    """RAPPORT_EVALUATION used to always require nettoyage/analyse, hard-
    failing any purely qualitative evaluation with no data file attached."""
    task, plan, result = _run_task(
        db_session, TaskType.RAPPORT_EVALUATION,
        "Évaluer l'impact du programme de cantines scolaires (évaluation qualitative)",
        {"objectif": 0.9, "donnees": 0.2},
    )
    assert "nettoyage" not in [s["type"] for s in plan.steps]
    assert result["final_status"] == "deliverable"


def test_rapport_evaluation_with_data_uses_data_driven_plan(db_session, pilot_xlsx):
    task, plan, result = _run_task(
        db_session, TaskType.RAPPORT_EVALUATION, "Évaluer la satisfaction socio-économique",
        {"objectif": 0.9, "contexte": 0.8, "donnees": 1.0, "livrables": 0.8,
         "contraintes": 0.7, "methode": 0.7, "ressources": 0.8},
        data_path=pilot_xlsx["path"],
    )
    assert "nettoyage" in [s["type"] for s in plan.steps]
    assert result["final_status"] == "deliverable"
    assert len(result["deliverables"]) == 2  # docx + pptx, same as analyse_donnees


def test_content_driven_document_lists_kb_sources_without_llm(db_session):
    """No LLM configured in tests - the document must honestly list what
    the Knowledge Base found rather than showing data-report boilerplate
    ('Aucune statistique disponible...') that doesn't apply to this task type."""
    kb = KnowledgeService(db_session)
    kb.add_item(
        title="Guide du Cadre Logique (GAR)", domain="Planification", content_type=KBContentType.METHOD,
        content="Le cadre logique structure objectifs, résultats, activités et indicateurs pour la planification de projet.",
    )
    task, plan, result = _run_task(
        db_session, TaskType.PLANIFICATION, "Planifier le projet avec un cadre logique", {"objectif": 0.9},
    )
    from docx import Document
    from src.services.artifact_service import ArtifactService
    artifact = ArtifactService(db_session).get_artifact(result["deliverables"][0]["artifact_id"])
    doc = Document(artifact.file_ref)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "Aucune statistique disponible" not in text
    assert "Guide du Cadre Logique" in text


def test_constraint_compliance_check_tolerates_plain_string_constraints(db_session):
    """ValidationService._check_constraint used to assume every constraint
    was a {"type", "value"} dict and crashed with AttributeError on a plain
    string - which nothing had exercised before hard_constraints actually
    had content in it (see test_reponse_ao_produces_real_deliverable...)."""
    from src.services.validation_service import ValidationService

    service = ValidationService(db_session)
    result = service.run_constraint_compliance_check(
        constraints={"hard_constraints": ["Délai de remise: 30 jours"], "soft_preferences": []},
        deliverable_metadata={"format": "docx"},
    )
    assert result["passed"] is True


def test_redaction_uses_uploaded_notes_as_raw_material_without_llm(db_session, tmp_path):
    """The core "externalisation de savoir acquis" (knowledge capitalization)
    use case: the user's own field notes get structured into the document,
    not silently ignored or replaced by an empty placeholder."""
    notes_path = tmp_path / "notes_terrain.txt"
    notes_path.write_text(
        "Constat: la maintenance des pompes s'essouffle après 18 mois faute de comité "
        "de gestion formé.\nLeçon apprise: former le comité de gestion dès la construction.",
        encoding="utf-8",
    )

    task, plan, result = _run_task(
        db_session, TaskType.REDACTION,
        "Rédiger une fiche de capitalisation à partir de mes notes de mission",
        {"objectif": 0.9, "contenu": 0.8}, data_path=str(notes_path),
    )
    assert result["final_status"] == "deliverable"

    from docx import Document
    from src.services.artifact_service import ArtifactService
    artifact = ArtifactService(db_session).get_artifact(result["deliverables"][0]["artifact_id"])
    doc = Document(artifact.file_ref)
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "comité de gestion" in text
    assert "reprise telle quelle" in text  # honest about no LLM having structured it
